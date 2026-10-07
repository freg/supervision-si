# -*- coding: utf-8 -*-
"""Rapport des anomalies réseau Nebula (livraison #703) -- logique PURE,
testée (nebula/tests/test_report.py) : lignes à plat depuis les anomalies
enrichies par les règles (rules/), exports CSV, Excel mis en forme (synthèse +
détail, couleurs par gravité, filtres, volets figés) et PDF (paysage, police
DejaVu si présente pour les flèches et tirets, sinon Helvetica avec
substitution des caractères hors Latin-1)."""
import csv
import io
import os
from datetime import datetime

SEV_ORDER = {"haute": 0, "moyenne": 1, "basse": 2, "info": 3}
SEV_LABEL = {"haute": "Haute", "moyenne": "Moyenne", "basse": "Basse", "info": "Info"}
SEV_COLOR = {"haute": "C62828", "moyenne": "EF6C00", "basse": "1565C0", "info": "757575"}
STATE_LABEL = {"validated": "validée", "applied": "appliquée", "failed": "échec", "hidden": "masquée", None: "ouverte"}
COLUMNS = [("site", "Site", 22), ("severity", "Gravité", 10), ("element", "Élément", 34), ("message", "Constat", 70),
           ("action", "Action proposée", 60), ("state", "État", 11), ("since", "Depuis", 17), ("kind", "Type", 24), ("rule_id", "Règle", 9)]
FONT_PATHS = ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans.ttf")


def _when(ts):
    return datetime.fromtimestamp(ts).strftime("%d/%m/%Y %H:%M") if ts else ""


def rows_from(items_by_site, site_names=None, first_seen=None, include_hidden=False):
    """{site_id: [anomalie enrichie]} -> lignes triées (gravité, site, élément). `first_seen` : {(site, id): ts}."""
    rows = []
    for sid, items in (items_by_site or {}).items():
        for a in items or []:
            if a.get("state") == "hidden" and not include_hidden:
                continue
            fs = (first_seen or {}).get((sid, a.get("id")))
            rows.append({"site_id": sid, "site": (site_names or {}).get(sid) or sid, "id": a.get("id"), "severity": a.get("severity") or "info",
                         "element": a.get("element") or "", "message": a.get("message") or "", "action": a.get("action") or "",
                         "state": STATE_LABEL.get(a.get("state"), a.get("state") or ""), "since": _when(fs), "since_ts": fs,
                         "kind": a.get("kind") or "", "rule_id": a.get("rule_id") or ""})
    rows.sort(key=lambda r: (SEV_ORDER.get(r["severity"], 9), r["site"], r["element"], r["message"]))
    return rows


def counts(rows):
    out = {k: 0 for k in SEV_ORDER}
    for r in rows:
        out[r["severity"]] = out.get(r["severity"], 0) + 1
    return out


def _cell(r, key):
    return SEV_LABEL.get(r["severity"], r["severity"]) if key == "severity" else r.get(key, "")


def to_csv(rows):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([label for _k, label, _w in COLUMNS])
    for r in rows:
        w.writerow([_cell(r, k) for k, _l, _w in COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(rows, title="Anomalies réseau", generated_at=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    generated_at = generated_at or datetime.now()
    wb = Workbook()
    syn = wb.active
    syn.title = "Synthèse"
    syn["A1"] = title
    syn["A1"].font = Font(size=14, bold=True)
    syn["A2"] = "Généré le %s" % generated_at.strftime("%d/%m/%Y à %H:%M")
    syn["A2"].font = Font(italic=True, color="666666")
    syn.append([])
    syn.append(["Gravité", "Nombre"])
    for c in syn[4]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="37474F")
    for sev, n in counts(rows).items():
        syn.append([SEV_LABEL[sev], n])
        syn.cell(row=syn.max_row, column=1).font = Font(bold=True, color=SEV_COLOR[sev])
    syn.append(["Total", len(rows)])
    syn.cell(row=syn.max_row, column=1).font = Font(bold=True)
    syn.append([])
    syn.append(["Site", "Haute", "Moyenne", "Basse", "Info"])
    hdr = syn.max_row
    for c in syn[hdr]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="37474F")
    by_site = {}
    for r in rows:
        by_site.setdefault(r["site"], {k: 0 for k in SEV_ORDER})[r["severity"]] += 1
    for site, cnt in sorted(by_site.items()):
        syn.append([site] + [cnt[k] for k in ("haute", "moyenne", "basse", "info")])
    syn.column_dimensions["A"].width = 32
    for col in "BCDE":
        syn.column_dimensions[col].width = 11

    ws = wb.create_sheet("Anomalies")
    ws.append([label for _k, label, _w in COLUMNS])
    for i, (_k, _l, width) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="37474F")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    sev_col = [k for k, _l, _w in COLUMNS].index("severity") + 1
    for r in rows:
        ws.append([_cell(r, k) for k, _l, _w in COLUMNS])
        n = ws.max_row
        for i in range(1, len(COLUMNS) + 1):
            ws.cell(row=n, column=i).alignment = Alignment(vertical="top", wrap_text=True)
        sc = ws.cell(row=n, column=sev_col)
        sc.font = Font(bold=True, color="FFFFFF")
        sc.fill = PatternFill("solid", fgColor=SEV_COLOR.get(r["severity"], "757575"))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(COLUMNS)), max(1, ws.max_row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_SUBST = {"↔": "<->", "→": "->", "←": "<-", "—": "-", "–": "-", "…": "...", "«": '"', "»": '"', "’": "'", "≥": ">=", "≤": "<="}


def _latin1(s):
    s = "".join(_SUBST.get(ch, ch) for ch in str(s))
    return s.encode("latin-1", "replace").decode("latin-1")


def to_pdf(rows, title="Anomalies réseau", generated_at=None, font_paths=FONT_PATHS):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    generated_at = generated_at or datetime.now()
    font, txt = "Helvetica", _latin1
    for p in font_paths or ():
        if os.path.exists(p):
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            try:
                pdfmetrics.registerFont(TTFont("DejaVuSans", p))
                bold = p.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")
                if os.path.exists(bold):
                    pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", bold))
                font, txt = "DejaVuSans", str
            except Exception:  # noqa: BLE001 -- police illisible : Helvetica
                pass
            break
    bold_font = font + "-Bold" if font == "DejaVuSans" else "Helvetica-Bold"
    try:
        from reportlab.pdfbase import pdfmetrics as _pm
        _pm.getFont(bold_font)
    except Exception:  # noqa: BLE001
        bold_font = font
    st = ParagraphStyle("c", fontName=font, fontSize=7.5, leading=9.2)
    st_h = ParagraphStyle("h", fontName=bold_font, fontSize=8, leading=9.5, textColor=colors.white)
    st_t = ParagraphStyle("t", fontName=bold_font, fontSize=15, leading=18)
    st_s = ParagraphStyle("s", fontName=font, fontSize=9, leading=11, textColor=colors.HexColor("#555555"))
    P = lambda s, style=st: Paragraph(escape(txt(s)), style)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
                            title=txt(title), author="supervision-si")
    c = counts(rows)
    story = [P(title, st_t), P("Généré le %s · %d anomalie(s) : %s" % (generated_at.strftime("%d/%m/%Y à %H:%M"), len(rows),
                                 ", ".join("%d %s" % (c[k], SEV_LABEL[k].lower()) for k in SEV_ORDER if c[k])), st_s), Spacer(1, 4 * mm)]
    cols = [("severity", 16), ("site", 30), ("element", 52), ("message", 92), ("action", 70), ("since", 17)]
    labels = {k: l for k, l, _w in COLUMNS}
    data = [[P(labels[k], st_h) for k, _w in cols]]
    st_sev = ParagraphStyle("sev", parent=st, textColor=colors.white, fontName=bold_font)
    for r in rows:
        data.append([P(_cell(r, k), st_sev if k == "severity" else st) for k, _w in cols])
    if not rows:
        data.append([P("Aucune anomalie.")] + [P("") for _ in cols[1:]])
    t = Table(data, colWidths=[w * mm for _k, w in cols], repeatRows=1)
    style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#37474F")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")])]
    for i, r in enumerate(rows, 1):
        style.append(("BACKGROUND", (0, i), (0, i), colors.HexColor("#" + SEV_COLOR.get(r["severity"], "757575"))))
        style.append(("TEXTCOLOR", (0, i), (0, i), colors.white))
    t.setStyle(TableStyle(style))
    story.append(t)
    doc.build(story)
    return buf.getvalue()


def text_summary(rows, title, max_rows=40):
    """Corps de courriel en texte : synthèse + lignes (les plus graves d'abord)."""
    c = counts(rows)
    out = [title, "", "%d anomalie(s) : %s" % (len(rows), ", ".join("%d %s" % (c[k], SEV_LABEL[k].lower()) for k in SEV_ORDER if c[k]) or "aucune"), ""]
    for r in rows[:max_rows]:
        out.append("[%s] %s — %s" % (SEV_LABEL.get(r["severity"], r["severity"]), r["site"], r["message"]))
        if r["action"]:
            out.append("    → %s" % r["action"])
        if r.get("since"):
            out.append("    depuis le %s" % r["since"])
    if len(rows) > max_rows:
        out.append("… et %d autre(s), voir le rapport joint ou la tuile Nebula." % (len(rows) - max_rows))
    return "\n".join(out)
