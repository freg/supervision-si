# -*- coding: utf-8 -*-
"""Tests #703 : rapport des anomalies (CSV, Excel, PDF) et alertes (récapitulatif, urgence)."""
import io
import os
import sys
import unittest
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "api"))
import alerting  # noqa: E402
import report  # noqa: E402

ITEMS = {"s1": [
    {"id": "a1", "kind": "ap_port_missing_ssid_vlan", "element": "VLAN 22 sur 16 port(s) de borne", "severity": "haute", "rule_id": "R-NET-06",
     "message": "VLAN 22 (SSID Campus) absent de 16 port(s) de borne : GS2220-50HP-1 port 2 (Borne) ↔ …", "action": "Compléter le profil « AP »", "state": None},
    {"id": "a2", "kind": "ap_port_all_vlans", "element": "16 port(s) de borne", "severity": "moyenne", "rule_id": "R-NET-07", "message": "16 port(s) en « All »", "action": "Liste explicite", "state": "validated"},
    {"id": "a3", "kind": "vlan_no_gateway", "element": "VLAN 10", "severity": "basse", "message": "VLAN 10 sans passerelle", "action": "", "state": "hidden"}],
    "s2": [{"id": "b1", "kind": "gateway_ip_not_host", "element": "VLAN 22", "severity": "haute", "message": "VLAN 22 : adresse du réseau", "action": "172.16.22.1/24", "state": None}]}
NAMES = {"s1": "Site Alpha", "s2": "Site Beta"}


class Rapport(unittest.TestCase):
    def setUp(self):
        self.rows = report.rows_from(ITEMS, NAMES, {("s1", "a1"): datetime(2026, 10, 7, 9, 30).timestamp()})

    def test_lignes(self):
        self.assertEqual([r["id"] for r in self.rows], ["a1", "b1", "a2"])          # masquée exclue, gravité puis site
        self.assertEqual((self.rows[0]["site"], self.rows[0]["since"], self.rows[2]["state"]), ("Site Alpha", "07/10/2026 09:30", "validée"))
        self.assertEqual(len(report.rows_from(ITEMS, NAMES, include_hidden=True)), 4)
        self.assertEqual(report.counts(self.rows), {"haute": 2, "moyenne": 1, "basse": 0, "info": 0})

    def test_csv(self):
        b = report.to_csv(self.rows)
        self.assertTrue(b.startswith("﻿".encode("utf-8")))
        lines = b.decode("utf-8-sig").splitlines()
        self.assertEqual(lines[0].split(";")[:3], ["Site", "Gravité", "Élément"])
        self.assertIn("Haute", lines[1])

    def test_excel(self):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(report.to_xlsx(self.rows, "Anomalies réseau", datetime(2026, 10, 7, 8, 0))))
        self.assertEqual(wb.sheetnames, ["Synthèse", "Anomalies"])
        syn = wb["Synthèse"]
        self.assertEqual(syn["A1"].value, "Anomalies réseau"); self.assertEqual((syn["A5"].value, syn["B5"].value), ("Haute", 2))
        ws = wb["Anomalies"]
        self.assertEqual(ws.max_row, 4); self.assertEqual(ws.freeze_panes, "A2"); self.assertTrue(ws.auto_filter.ref.startswith("A1:"))
        self.assertEqual(ws["B2"].value, "Haute"); self.assertEqual(ws["B2"].fill.fgColor.rgb[-6:], "C62828")

    def test_pdf(self):
        pdf = report.to_pdf(self.rows, "Anomalies réseau")
        self.assertTrue(pdf.startswith(b"%PDF")); self.assertGreater(len(pdf), 1500)
        sans = report.to_pdf(self.rows, "Anomalies", font_paths=())                 # Helvetica : caractères hors Latin-1 remplacés
        self.assertTrue(sans.startswith(b"%PDF"))
        self.assertTrue(report.to_pdf([], "Vide").startswith(b"%PDF"))
        self.assertEqual(report._latin1("a ↔ b — c…"), "a <-> b - c...")

    def test_texte(self):
        t = report.text_summary(self.rows, "Récapitulatif")
        self.assertIn("3 anomalie(s) : 2 haute, 1 moyenne", t); self.assertIn("[Haute] Site Alpha", t); self.assertIn("→ Compléter le profil", t)


class Alertes(unittest.TestCase):
    def test_reglages(self):
        c = alerting.normalize({"recipients": "a@exemple.fr; mauvais , B@Exemple.fr", "daily": {"enabled": True, "time": "25:00", "weekdays": [1, 9, "3"], "formats": ["pdf", "doc"]},
                                "urgent": {"delay_minutes": 99999, "min_severity": "critique"}, "check_minutes": 1})
        self.assertEqual(c["recipients"], ["a@exemple.fr", "b@exemple.fr"])
        self.assertEqual((c["daily"]["time"], c["daily"]["weekdays"], c["daily"]["formats"]), ("08:00", [1, 3], ["pdf"]))
        self.assertEqual((c["urgent"]["delay_minutes"], c["urgent"]["min_severity"], c["check_minutes"]), (1440, "haute", 5))
        self.assertEqual(alerting.invalid_recipients("a@exemple.fr; mauvais"), ["mauvais"])
        self.assertEqual(alerting.normalize(None), alerting.normalize({}))

    def test_urgence_apres_delai_sans_retour_a_la_normale(self):
        cfg = alerting.normalize({"urgent": {"enabled": True, "delay_minutes": 15, "min_severity": "haute"}})["urgent"]
        rows = report.rows_from(ITEMS, NAMES)
        st, urg, rec = alerting.track({}, rows, 1000, cfg)
        self.assertEqual((urg, rec), ([], []))                                        # apparue : on attend
        st, urg, _ = alerting.track(st, rows, 1000 + 14 * 60, cfg)
        self.assertEqual(urg, [])
        st, urg, _ = alerting.track(st, rows, 1000 + 15 * 60, cfg)
        self.assertEqual(sorted(u["id"] for u in urg), ["a1", "b1"])                 # haute seulement, pas a2 (moyenne)
        st, urg, _ = alerting.track(st, rows, 1000 + 16 * 60, cfg)
        self.assertEqual(urg, [])                                                     # une seule fois
        # a1 corrigée : retour à la normale ; b1 toujours là
        st, urg, rec = alerting.track(st, [r for r in rows if r["id"] != "a1"], 1000 + 20 * 60, cfg)
        self.assertEqual([r["id"] for r in rec], ["a1"])
        self.assertEqual([e["id"] for e in alerting.resolved_since(st, 1000)], ["a1"])
        # anomalie brève (disparue avant le délai) : ni alerte ni retour à la normale
        st2, _u, _r = alerting.track({}, rows, 0, cfg)
        st2, u2, r2 = alerting.track(st2, [], 5 * 60, cfg)
        self.assertEqual((u2, r2), ([], []))
        # rappel toutes les 2 h
        cfg_r = dict(cfg, repeat_hours=2, delay_minutes=0)
        s3, u3, _ = alerting.track({}, rows, 0, cfg_r)
        s3, u3b, _ = alerting.track(s3, rows, 3600, cfg_r)
        s3, u3c, _ = alerting.track(s3, rows, 7200, cfg_r)
        self.assertEqual((len(u3), len(u3b), len(u3c)), (2, 0, 2))
        self.assertEqual(alerting.track({}, rows, 0, dict(cfg, enabled=False))[1], [])

    def test_recapitulatif(self):
        d = alerting.normalize({"daily": {"enabled": True, "time": "08:00", "weekdays": [1, 2, 3, 4, 5]}})["daily"]
        mer = datetime(2026, 10, 7, 8, 5)                                             # mercredi
        self.assertTrue(alerting.daily_due(d, None, mer))
        self.assertFalse(alerting.daily_due(d, datetime(2026, 10, 7, 8, 1).timestamp(), mer))   # déjà envoyé
        self.assertTrue(alerting.daily_due(d, datetime(2026, 10, 6, 8, 1).timestamp(), mer))    # envoyé hier
        self.assertFalse(alerting.daily_due(d, None, datetime(2026, 10, 7, 7, 59)))
        self.assertFalse(alerting.daily_due(d, None, datetime(2026, 10, 7, 15, 0)))              # rattrapage borné à 6 h
        self.assertFalse(alerting.daily_due(d, None, datetime(2026, 10, 10, 8, 5)))              # samedi
        self.assertEqual(alerting.next_daily(d, datetime(2026, 10, 9, 9, 0)), datetime(2026, 10, 12, 8, 0))  # vendredi -> lundi
        self.assertIsNone(alerting.next_daily(dict(d, enabled=False), mer))


if __name__ == "__main__":
    unittest.main()
