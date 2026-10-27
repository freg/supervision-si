# -*- coding: utf-8 -*-
"""Tests #878 : temps estimés par action (marge ×3), regroupements de dépendances, phases, calendrier, Excel et page."""
import datetime as dt
import io
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import etat_plan as ep  # noqa: E402


def L(key, niveau, actions, famille="debian", version="9", type_="CT", dependants=(), roles=(), score=5, name=None):
    return {"key": key, "name": name or key, "type": type_, "niveau": niveau, "actions": list(actions), "score": score, "dependants": list(dependants),
            "roles": list(roles), "os": {"famille": famille, "version": version, "nom": "%s %s" % (famille, version)}, "libelle": None}


ETAT = {"systemes": [
    L("sql2", "urgent", ["mettre à niveau le système (Debian 9 → version maintenue)", "appliquer les mises à jour de sécurité", "remplacer mysql 5.5 par une version maintenue"],
      dependants=["tickets", "abos"], roles=["base de données"], score=11, name="sql2.intranet"),
    L("tickets", "à planifier", ["remplacer php 7.3 par une version maintenue"], version="12", roles=["site ou application web"], score=4),
    L("abos", "à surveiller", ["mettre à jour les paquets"], version="12", score=1),
    L("seul", "à planifier", ["mettre à niveau le système (Debian 11 → version maintenue)"], version="11", score=3),
    L("vieux", "urgent", ["mettre à niveau le système (Debian 4 → version maintenue)", "limiter l'accès depuis Internet en attendant la mise à niveau"], version="4", score=9),
    L("arret", "inactif", ["archiver puis supprimer s'il n'est plus utile"], score=0),
    L("ok1", "ok", [], version="13", score=0)],
    "dependances": [{"de": "tickets", "vers": "sql2", "port": 3306, "n": 4, "nature": "base de données"}, {"de": "abos", "vers": "sql2", "port": 3306, "n": 2, "nature": "base de données"}]}
LUNDI = dt.date(2026, 11, 2)


class Plan(unittest.TestCase):
    def test_heures(self):
        l = ETAT["systemes"][0]
        h, _ = ep.heures_action("mettre à niveau le système (Debian 9 → version maintenue)", l)
        self.assertEqual(h, 16.0 + 3)                                  # écart 4 versions -> reconstruction + rôle base de données
        self.assertEqual(ep.heures_action("remplacer mysql 5.5 par une version maintenue", l, mise_a_niveau_os=True)[0], 3.0)
        self.assertEqual(ep.heures_action("mettre à niveau le système (Debian 11 → …)", ETAT["systemes"][3])[0], 4.0 + 3 * 2)
        t = ep.taches(l, marge=3)
        self.assertEqual([x["heures"] for x in t], [57.0, 4.5, 9.0])  # brut × 3 ; mysql à moitié (fait avec le système)

    def test_groupes(self):
        g = {x["pivot"]["key"]: [m["key"] for m in x["membres"]] for x in ep.groupes(ETAT)}
        self.assertEqual(g["sql2"], ["sql2", "tickets", "abos"])       # pivot d'abord, puis par score
        self.assertNotIn("ok1", g)

    def test_plan_phases_et_calendrier(self):
        p = ep.plan(ETAT, marge=3, personnes=1, debut=LUNDI)
        ph = {x["phase"]: x for x in p["phases"]}
        self.assertEqual(sorted(ph), ["1", "2", "3", "5"])
        t1 = [t for t in p["taches"] if t["phase"] == "1"]
        self.assertEqual({t["key"] for t in t1}, {"sql2", "vieux"})    # sécurité / exposition d'abord
        t2 = [t for t in p["taches"] if t["phase"] == "2"]
        self.assertTrue(any(t["groupe"].startswith("sql2.intranet et 2") for t in t2))
        ordre = [t["key"] for t in t2 if t["groupe"].startswith("sql2")]
        self.assertEqual(ordre[0], "sql2")
        self.assertEqual(p["taches"][0]["debut"], "2026-11-02")
        for t in p["taches"]:
            self.assertLess(dt.date.fromisoformat(t["fin"]).weekday(), 5)
            self.assertLessEqual(t["debut"], t["fin"])
        self.assertEqual(p["total"]["heures"], round(3 * p["total"]["brut"], 1))
        p2 = ep.plan(ETAT, marge=3, personnes=2, debut=LUNDI)
        self.assertLess(p2["total"]["fin"], p["total"]["fin"])        # deux personnes : fin plus tôt
        self.assertEqual(ep.par_systeme(p)["sql2"]["heures"], 70.5)

    def test_excel_et_page(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl absent")
        import etat
        import etat_excel
        import etat_rapport
        e = etat.assemble({}, maintenant=1793000000)
        e["systemes"], e["dependances"] = ETAT["systemes"], ETAT["dependances"]
        for l in e["systemes"]:
            l.update(status="running", composants=[], raisons=[], ips=[], maj_at=None, attente_secu=None, usage="moyen", expose=False,
                     os=dict(l["os"], statut="obsolète", fin="2022-06-30"))
        p = ep.plan(e, debut=LUNDI)
        wb = openpyxl.load_workbook(io.BytesIO(etat_excel.classeur(e, maintenant=1793000000, plan=p)))
        self.assertEqual(wb.sheetnames[1], "Planning")
        ws = wb["Planning"]
        self.assertIn("multipliée par 3", ws["A2"].value)
        self.assertTrue(any("sql2.intranet et 2" in str(c.value) for c in ws["A"]))
        sys_ = wb["Systèmes"]
        self.assertEqual(sys_.cell(row=1, column=14).value, "Temps estimé (jours, marge comprise)")
        html = etat_rapport.page(e, plan=p)
        self.assertIn("Planning estimé", html)
        self.assertIn("sql2.intranet et 2", html)


if __name__ == "__main__":
    unittest.main()
