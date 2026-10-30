# -*- coding: utf-8 -*-
"""Tests #873 : audit réseau dans l'état des lieux (système déduit des services, hôtes découverts, complément des connus, Excel)."""
import os
import sys
import time
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import etat  # noqa: E402

NOW = time.mktime((2026, 10, 27, 12, 0, 0, 0, 0, 0))
NMS = {"ip": "192.0.2.20", "mac": "02:00:00:00:00:20", "vendor": "Dell", "hostname": "nms.exemple",
       "os": {"name": "Linux 2.6.18 - 2.6.32", "accuracy": 95, "family": "Linux"},
       "ports": [{"port": 22, "service": "ssh", "product": "OpenSSH", "version": "5.1p1 Debian 5", "extrainfo": "protocol 2.0 lenny"},
                 {"port": 80, "service": "http", "product": "Apache httpd", "version": "2.2.9", "extrainfo": "(Debian) PHP/5.2.6-1+lenny16"},
                 {"port": 3306, "service": "mysql", "product": "MySQL", "version": "5.0.51a-24+lenny5"}]}
WIN = {"ip": "192.0.2.100", "os": {"name": "Microsoft Windows Server 2012 R2", "accuracy": 96, "family": "Windows"},
       "ports": [{"port": 3389, "service": "ms-wbt-server", "product": "Microsoft Terminal Services"}]}
IMPR = {"ip": "192.0.2.150", "vendor": "Exemple Print", "os": None, "ports": [{"port": 9100, "service": "jetdirect"}]}


class AuditReseau(unittest.TestCase):
    def test_systeme_deduit(self):
        s = etat.systeme_audit(NMS)
        self.assertEqual((s["os"]["id"], s["os"]["version"]), ("debian", "5"))
        self.assertEqual({a["nom"]: a["version"] for a in s["apps"]}, {"openssh": "5.1", "apache-httpd": "2.2", "php": "5.2", "mysql": "5.0"})
        w = etat.systeme_audit(WIN)["os"]
        self.assertEqual((w["id"], w["version"]), ("windows-server", "2012r2"))
        self.assertIsNone(etat.systeme_audit(IMPR)["os"].get("id"))
        deb12 = etat.systeme_audit({"ip": "192.0.2.9", "ports": [{"port": 22, "product": "OpenSSH", "version": "9.2p1 Debian 2+deb12u3"}]})
        self.assertEqual(deb12["os"]["version"], "12")

    def test_hotes_audit_le_plus_recent(self):
        r = [("a1", "h", "2026-10-20T00:00:00Z", {"hosts": [dict(NMS, hostname="ancien")]}),
             ("a2", "h", "2026-10-27T00:00:00Z", {"hosts": [NMS, WIN]})]
        h = etat.hotes_audit(r)
        self.assertEqual(h["192.0.2.20"]["hostname"], "nms.exemple")
        self.assertEqual(sorted(h), ["192.0.2.100", "192.0.2.20"])

    def test_assemble(self):
        autres = [{"name": "optick3", "role": "serveur", "ok": True, "systeme": {"os": {"id": "debian", "version": "4", "nom": "Debian 4.0"},
                   "ips": ["192.0.2.30"], "ecoute": [], "paquets": {}, "apps": []}}]
        scan = etat.hotes_audit([("a", "h", "2026-10-27T00:00:00Z", {"hosts": [NMS, WIN, IMPR, dict(NMS, ip="192.0.2.30", hostname="optick3")]})])
        e = etat.assemble({}, maintenant=NOW, autres=autres, scan=scan)
        par = {l["key"]: l for l in e["systemes"]}
        self.assertIn("hote:optick3", par)                                  # connu : complété, pas dupliqué
        self.assertEqual(par["hote:optick3"]["source"], "relevé")
        self.assertEqual(len(par["hote:optick3"]["reseau"]), 3)
        self.assertNotIn("ip:192.0.2.30", par)
        nms = par["ip:192.0.2.20"]
        self.assertEqual((nms["type"], nms["source"], nms["os"]["statut"]), ("découvert", "audit réseau", "obsolète"))
        self.assertEqual(nms["niveau"], "urgent")
        self.assertIn("apache-httpd", [c["nom"] for c in nms["composants"] if c["statut"] == "obsolète"])
        self.assertEqual(par["ip:192.0.2.150"]["os"]["statut"], "inconnu")
        self.assertIn("identifier l'hôte", " ".join(par["ip:192.0.2.150"]["actions"]))
        self.assertEqual(e["resume"]["decouverts"], 3)
        self.assertEqual(e["resume"]["audit_reseau"], 4)

    def test_excel(self):
        try:
            import openpyxl  # noqa: F401
            import etat_excel
        except ImportError:
            self.skipTest("openpyxl absent")
        import io
        scan = etat.hotes_audit([("a", "h", "2026-10-27T00:00:00Z", {"hosts": [NMS, WIN]})])
        e = etat.assemble({}, maintenant=NOW, scan=scan)
        fn = getattr(etat_excel, "classeur", None) or getattr(etat_excel, "build", None)
        if fn is None:
            self.skipTest("point d'entrée Excel non trouvé")
        data = fn(e) if fn.__code__.co_argcount == 1 else fn(e, maintenant=NOW)
        wb = openpyxl.load_workbook(io.BytesIO(data if isinstance(data, bytes) else data.getvalue()))
        self.assertIn("Audit réseau", wb.sheetnames)
        self.assertEqual(wb["Audit réseau"].max_row, 5)


class HotesRelayes(unittest.TestCase):
    def test_hotes_relayes(self):
        rows = [("super-hub", "super", "2026-10-27T10:20:06Z", {"systemes": [
            {"name": "super", "zone": "local", "ok": True, "systeme": {}},
            {"name": "optick3", "zone": "distant", "role": "serveur", "ok": True, "systeme": {"os": {"nom": "Debian GNU/Linux 4.0"}, "noyau": "2.6.18-6-686",
             "ips": ["192.0.2.30"], "ecoute": [{"port": 80, "hote": "192.0.2.30"}, {"port": 25, "hote": "127.0.0.1"}, {"port": 22, "hote": "*"}]}},
            {"name": "nms", "zone": "distant", "ok": False, "error": "ssh refusé", "systeme": None}]})]
        h = etat.hotes_relayes(rows)
        self.assertEqual([x["name"] for x in h], ["nms", "optick3"])
        self.assertEqual(h[1]["ports"], [22, 80])
        self.assertEqual((h[0]["ok"], h[0]["error"], h[0]["os"]), (False, "ssh refusé", None))
        self.assertEqual(etat.hotes_relayes([]), [])

    def test_hotes_relayes_nas_backuppc_945(self):
        ssh = [("super", "super", "t1", {"nas": [{"name": "nas-grpi-1", "host": "192.0.2.92", "ok": True, "systeme": {"version": "5.2.9", "modele": "TS-883XU-RP"}}]})]
        snmp = [("super", "super", "t2", {"nas": [{"name": "NAS-GRPI-1", "host": "192.0.2.92", "ok": True, "systeme": {"modele": "TS-883XU-RP"}}, {"name": "nas-2", "host": "192.0.2.93", "ok": False, "error": "timeout", "systeme": {}}]})]
        h = etat.hotes_relayes_nas(ssh, snmp)
        self.assertEqual([(x["name"], x["via"], x["ok"], x["vue"]) for x in h], [("nas-2", "SNMP", False, "nas"), ("nas-grpi-1", "SSH + SNMP", True, "sauvegardes")])
        self.assertEqual(h[1]["os"], "QTS 5.2.9")
        self.assertEqual(etat.hotes_relayes_nas([], snmp)[1]["via"], "SNMP")
        b = etat.hotes_relayes_backuppc([("super", "super", "t", {"serveur": {"nom": "fichiers", "relais": True, "version": "3.2.1"}, "summary": {"hotes": 4, "echecs": 2}}),
                                         ("fichiers", "fichiers", "t", {"serveur": {"version": "3.2.1"}})])
        self.assertEqual(len(b), 1); self.assertEqual((b[0]["name"], b[0]["os"], b[0]["noyau"]), ("fichiers", "BackupPC 3.2.1", "4 hôte(s), 2 échec(s)"))


if __name__ == "__main__":
    unittest.main()
