# -*- coding: utf-8 -*-
"""#857 : état des lieux -- fin de support, composants, rôles, usage, exposition, dépendances, niveaux, Excel."""
import io
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import etat as E  # noqa: E402
import etat_excel as X  # noqa: E402
import etat_rapport as RP  # noqa: E402

T = E._date("2026-10-26")
J = 86400


def sy(os_id, ver, nom, **kw):
    return dict({"os": {"id": os_id, "version": ver, "nom": nom}, "paquets": {}, "apps": [], "ecoute": [], "cnx_in": [], "cnx_out": [], "services": [], "ips": []}, **kw)


NODES = {"pve-2": {"pveversion": "pve-manager/5.4-15/d0ec33c6 (running kernel: 4.15.18-30-pve)",
                   "systeme": sy("debian", "9", "Debian GNU/Linux 9 (stretch)", paquets={"pve-manager": "5.4-15"}, maj_at=T - 400 * J, attente=10, attente_secu=0,
                                 ips=["198.51.100.2"], ecoute=[{"port": 8006, "hote": "0.0.0.0", "proc": "pveproxy"}]),
                   "guests": [
                       {"vmid": 100, "type": "lxc", "status": "running", "name": "THD85", "conf": {"ips": ["198.51.100.7"]},
                        "charge": {"cpu_avg_pct": 12, "cpu_max_pct": 80, "mem_max_pct": 70},
                        "systeme": sy("debian", "9", "Debian GNU/Linux 9 (stretch)", maj_at=T - 700 * J, attente=40, attente_secu=3,
                                      paquets={"php7.0-common": "7.0.33", "apache2": "2.4.25"}, apps=[{"nom": "owncloud", "version": "9.1.8"}],
                                      ecoute=[{"port": 443, "hote": "0.0.0.0", "proc": "apache2"}, {"port": 22, "hote": "0.0.0.0", "proc": "sshd"},
                                              {"port": 3306, "hote": "127.0.0.1", "proc": "mysqld"}],
                                      cnx_in=[{"port": 443, "n": 40, "clients": 25}], cnx_out=[{"hote": "198.51.100.12", "port": 5432, "n": 2}], requetes=20000)},
                       {"vmid": 112, "type": "lxc", "status": "running", "name": "SOG-PM", "conf": {"ips": ["198.51.100.12"]},
                        "systeme": sy("debian", "12", "Debian GNU/Linux 12 (bookworm)", maj_at=T - 3 * J, attente=0, attente_secu=0,
                                      paquets={"postgresql-15": "15.8", "php8.2-common": "8.2.20"},
                                      ecoute=[{"port": 5432, "hote": "0.0.0.0", "proc": "postgres"}])},
                       {"vmid": 102, "type": "lxc", "status": "stopped", "name": "MED10", "conf": {"ips": []},
                        "systeme": sy("debian", "8", "Debian GNU/Linux 8 (jessie)", maj_at=T - 900 * J)},
                       {"vmid": 105, "type": "qemu", "status": "running", "name": "ATHENA", "conf": {}, "systeme": None},
                       {"vmid": 999, "type": "lxc", "status": "stopped", "name": "modèle", "template": True}]},
         "pve-1": {"ok": False}}


class Etat(unittest.TestCase):
    def test_fin_support(self):
        self.assertEqual(E.fin_support("debian", "12"), ("2028-06-30", "12"))
        self.assertEqual(E.fin_support("php", "8.1"), ("2025-12-31", "8.1"))
        self.assertEqual(E.fin_support("ubuntu", "22.04"), ("2027-06-30", "22.04"))
        self.assertEqual(E.fin_support("owncloud", "9.1.8")[0], "2017-12-31")
        self.assertEqual(E.statut_fin("2026-12-31", T), "fin proche")
        self.assertEqual(E.statut_fin(None, T), "inconnu")
        self.assertTrue(E.public("198.51.100.7") is False)          # plage de documentation : réservée
        self.assertTrue(E.public("54.36.82.190"))
        self.assertFalse(E.public("192.168.0.4"))

    def test_assemble(self):
        r = E.assemble(NODES, maintenant=T, libelles={"pve-2/100": "Partage de fichiers THD"})
        L = {l["key"]: l for l in r["systemes"]}
        self.assertNotIn("pve-2/999", L)
        thd = L["pve-2/100"]
        self.assertEqual(thd["os"]["statut"], "obsolète")
        self.assertEqual(thd["niveau"], "urgent")
        self.assertEqual(thd["libelle"], "Partage de fichiers THD")
        self.assertEqual({c["nom"]: c["statut"] for c in thd["composants"]}, {"php": "obsolète", "owncloud": "obsolète"})
        self.assertIn("partage de fichiers (ownCloud)", thd["roles"])
        self.assertNotIn("base de données", thd["roles"])                         # 3306 sur 127.0.0.1 : interne
        self.assertEqual(thd["usage"], "fort")
        self.assertEqual(thd["utilise"], ["pve-2/112"])
        sog = L["pve-2/112"]
        self.assertEqual(sog["dependants"], ["pve-2/100"])
        self.assertEqual(sog["niveau"], "à surveiller")                           # PHP 8.2 : fin de support proche
        self.assertEqual({c["nom"]: c["statut"] for c in sog["composants"]}, {"php": "fin proche", "postgresql": "supporté"})
        self.assertEqual(L["pve-2/102"]["niveau"], "inactif")
        self.assertEqual(L["pve-2/102"]["actions"], ["archiver puis supprimer s'il n'est plus utile"])
        self.assertEqual((L["pve-2/105"]["os"]["statut"], L["pve-2/105"]["niveau"]), ("non relevé", "à surveiller"))
        noeud = L["node:pve-2"]
        self.assertEqual(noeud["os"]["statut"], "obsolète")
        self.assertIn("proxmox-ve", [c["nom"] for c in noeud["composants"]])
        self.assertEqual(r["dependances"][0]["nature"], "base de données")
        self.assertEqual(r["systemes"][0]["niveau"], "urgent")
        self.assertEqual(r["resume"]["inactif"], 1)

    def test_roles_noeud_et_declaration(self):
        sy_ = {"paquets": {"postfix": "3.7", "pve-manager": "9.0"}, "ecoute": [{"port": 3128, "hote": "0.0.0.0", "proc": "spiceproxy"},
                                                                         {"port": 25, "hote": "127.0.0.1", "proc": "master"},
                                                                         {"port": 8006, "hote": "0.0.0.0", "proc": "pveproxy"}]}
        self.assertEqual(E.roles(sy_), ["hyperviseur Proxmox"])                     # ni « proxy » ni « messagerie »
        r = E.assemble(NODES, maintenant=T, manuels={"pve-2/105": {"famille": "windows-server", "version": "2016", "maj": "2026-10-01"}})
        L = {l["key"]: l for l in r["systemes"]}
        a = L["pve-2/105"]
        self.assertEqual((a["source"], a["os"]["fin"], a["os"]["statut"], a["maj_jours"]), ("déclaré", "2027-01-12", "fin proche", 25))   # fin < 6 mois
        self.assertEqual(r["resume"]["inconnus"], 0)
        r = E.assemble(NODES, maintenant=T)
        self.assertEqual(r["resume"]["inconnus"], 1)
        self.assertIn("déclarer le système", " ".join(next(l for l in r["systemes"] if l["key"] == "pve-2/105")["actions"]))

    def test_lu_hors_ligne(self):
        lus = E.hors_ligne([("a", "h", "2026-10-26T10:00:00", {"nodes": [{"name": "pve-2", "guests": [
            {"vmid": 105, "systeme": {"os": {"id": "debian", "version": "8", "nom": "Debian 8"}, "maj_at": T - 3000 * J, "paquets": {}, "apps": [],
                                      "services": ["apache2"], "ips": ["198.51.100.15"], "ecoute": [], "cnx_out": [], "cnx_in": []}},
            {"vmid": 100, "systeme": None}]}]})])
        self.assertEqual(sorted(lus), ["pve-2/105"])
        r = E.assemble(NODES, maintenant=T, lus=lus, manuels={"pve-2/105": {"famille": "windows", "version": "11"}})
        a = next(l for l in r["systemes"] if l["key"] == "pve-2/105")
        self.assertEqual((a["source"], a["os"]["statut"], a["services"]), ("relevé (activité)", "obsolète", ["apache2"]))   # le relevé prime sur la déclaration
        self.assertIn("198.51.100.15", a["ips"])

    def test_audit_html(self):
        r = E.assemble(NODES, maintenant=T, libelles={"pve-2/100": "Partage <THD>"})
        el = RP.elements(r, T)
        self.assertEqual(el[0]["rang"], 0)                                             # échéances dépassées d'abord
        self.assertNotIn("MED10", " ".join(e["systeme"] for e in el))                  # arrêtés exclus
        self.assertTrue(any(e["quoi"].startswith("ownCloud") for e in el))
        self.assertEqual(RP.echeance("2026-12-31", T)[1], "orange")
        self.assertEqual(RP.echeance("2020-06-30", T)[0], "dépassée depuis 6 an(s)")
        h = RP.page(r, "Organisation", T)
        self.assertIn("Partage &lt;THD&gt;", h)                                         # échappé
        self.assertIn("Calendrier des fins de support", h)
        self.assertIn("window.print()", h)

    def test_si_interne_et_non_suivis(self):
        rep = [("ag", "h", "2026-10-26T10:00:00", {"systemes": [
            {"name": "nms", "role": "serveur", "ok": True, "systeme": sy("debian", "9", "Debian 9", ips=["192.0.2.30"], maj_at=T - 10 * J,
                                                                        cnx_out=[{"hote": "198.51.100.12", "port": 5432, "n": 1}])},
            {"name": "pc-alice", "role": "poste", "ok": True, "systeme": sy("windows", "10", "Windows 10 Pro", ips=["192.0.2.101"],
                                                                            cnx_out=[{"hote": "192.0.2.30", "port": 443, "n": 2}, {"hote": "192.0.2.50", "port": 445, "n": 4}])},
            {"name": "thd-agent", "role": "serveur", "ok": True, "systeme": sy("debian", "9", "Debian 9", ips=["198.51.100.7"])},
            {"name": "nas", "role": "serveur", "ok": False, "error": "Permission denied", "systeme": None}]})]
        r = E.assemble(NODES, maintenant=T, autres=E.autres_systemes(rep))
        L = {l["key"]: l for l in r["systemes"]}
        self.assertNotIn("hote:thd-agent", L)                                          # même adresse que pve-2/100 : pas de doublon
        self.assertEqual((L["hote:nms"]["type"], L["hote:nms"]["os"]["statut"]), ("serveur", "obsolète"))
        self.assertEqual((L["hote:pc-alice"]["type"], L["hote:pc-alice"]["os"]["fin"]), ("poste", "2025-10-14"))
        self.assertEqual(L["hote:nas"]["niveau"], "à surveiller")
        self.assertIn("pve-2/112", L["hote:nms"]["utilise"])
        self.assertIn("hote:pc-alice", L["hote:nms"]["dependants"])
        ns = {x["ip"]: x for x in r["non_suivis"]}
        self.assertEqual((ns["192.0.2.50"]["postes"], ns["192.0.2.50"]["natures"], ns["192.0.2.50"]["interne"]), (1, ["partage de fichiers"], True))

    def test_php_maintenu_par_la_distribution(self):
        c = E.composants({"paquets": {"php7.4-common": "7.4.33"}}, "supporté", T)
        self.assertEqual(c[0]["statut"], "maintenu par la distribution")

    def test_tendance_et_excel(self):
        r = E.assemble(NODES, maintenant=T)
        h = [("2026-10-25", E.instantane(r)), ("2026-10-24", {"resume": {"urgent": 4}})]
        t = E.tendance(h)
        self.assertEqual([x["jour"] for x in t], ["2026-10-24", "2026-10-25"])
        data = X.classeur(r, t, "Organisation")
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Synthèse", "Systèmes", "Dépendances", "Évolution"])
        texte = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
        self.assertIn("ne reçoivent plus de correctifs", texte)
        self.assertNotIn("3306", texte)                                            # pas de jargon : pas de ports
        self.assertNotIn("php", texte.lower())                                     # ni de noms de paquets
        self.assertEqual(wb["Systèmes"].max_row, len(r["systemes"]) + 1)
        sog = next(row for row in wb["Systèmes"].iter_rows(values_only=True) if row[0] == "SOG-PM (pve-2/112)")
        self.assertEqual(sog[12], "Planifier la mise à jour d'un logiciel installé avant le 31/12/2026")
        self.assertIn("support se termine bientôt", sog[11])
        self.assertEqual(wb["Évolution"]["A2"].value, "24/10/2026")


if __name__ == "__main__":
    unittest.main()
