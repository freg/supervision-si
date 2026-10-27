# -*- coding: utf-8 -*-
"""#857 : état des lieux des systèmes -> classeur Excel pour les décideurs : court, en clair, sans jargon.
Feuilles : Synthèse (chiffres clés, priorités), Systèmes (un par ligne, niveau et action proposée), Dépendances,
Évolution (suivi dans le temps). Testé dans test_z_etat.py."""
import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TEINTES = {"urgent": "F4B6B6", "à planifier": "F9D9A8", "à surveiller": "FFF2B3", "ok": "CDEBC5", "inactif": "E0E0E0"}
LIBELLE_NIVEAU = {"urgent": "Urgent", "à planifier": "À planifier", "à surveiller": "À surveiller", "ok": "À jour", "inactif": "Arrêté"}
TYPE = {"nœud": "Serveur physique", "VM": "Machine virtuelle", "CT": "Conteneur", "serveur": "Serveur", "poste": "Poste de travail", "découvert": "Appareil découvert sur le réseau"}
USAGE = {"fort": "Forte", "moyen": "Moyenne", "faible": "Faible", "nul": "Arrêté"}
TITRE = Font(bold=True, size=14)
GRAS = Font(bold=True)
ENTETE = PatternFill("solid", fgColor="D9E1F2")


def _n(x):
    """3.0 -> « 3 » ; 1.5 -> « 1,5 » (#882)."""
    try:
        f = float(x)
    except (TypeError, ValueError):
        return str(x)
    return str(int(f)) if f == int(f) else ("%g" % f).replace(".", ",")


def _jour(t):
    return dt.datetime.fromtimestamp(t, dt.timezone.utc).strftime("%d/%m/%Y") if t else "—"


def _fin(d):
    try:
        return dt.datetime.strptime(d, "%Y-%m-%d").strftime("%d/%m/%Y")
    except (TypeError, ValueError):
        return "—"


def nom(l):
    n = l.get("libelle") or ("%s (%s)" % (l["name"], l["key"]) if l["type"] != "nœud" else "Serveur %s" % l["name"])
    v = l.get("variante_de")   # #879 : copie, clone, préparation, test -- listés en fin de tableau
    return "%s — %s de %s" % (n, v["role"], v["name"]) if v else n


def raisons_claires(l):
    """Raisons sans jargon technique (pas de numéros de port, pas de noms de paquets)."""
    out = []
    if l["os"]["statut"] == "obsolète":
        out.append("système d'exploitation qui ne reçoit plus de correctifs depuis le %s" % _fin(l["os"]["fin"]))
    elif l["os"]["statut"] == "fin proche":
        out.append("fin des correctifs le %s" % _fin(l["os"]["fin"]))
    if l.get("attente_secu"):
        out.append("correctifs de sécurité disponibles mais non appliqués")
    if l.get("maj_jours") is not None and l["maj_jours"] > 365:
        out.append("pas mis à jour depuis plus d'un an")
    if any(c["statut"] == "obsolète" for c in l["composants"]):
        out.append("logiciel installé dans une version qui n'est plus maintenue")
    elif any(c["statut"] == "fin proche" for c in l["composants"]):
        out.append("logiciel installé dont le support se termine bientôt")
    if l["os"]["statut"] in ("inconnu", "non relevé") and l["status"] == "running":
        out.append("système non identifié automatiquement")
    if l.get("expose") and l["score"] >= 3:
        out.append("accessible depuis Internet")
    if l["dependants"]:
        out.append("d'autres systèmes en dépendent")
    return "; ".join(out) or "—"


def _clair(t):
    """Action sans nom de logiciel ni numéro de version technique."""
    import re
    m = re.match(r"^planifier le remplacement de .* avant le (\d{4}-\d{2}-\d{2})$", t)
    if m:
        return "planifier la mise à jour d'un logiciel installé avant le %s" % _fin(m.group(1))
    if t.startswith("remplacer "):
        return "remplacer un logiciel installé par une version maintenue"
    if t.startswith("mettre à niveau le système"):
        return "mettre à niveau le système d'exploitation vers une version maintenue"
    if t.startswith("planifier la mise à niveau du système avant le "):
        return "planifier la mise à niveau du système avant le %s" % _fin(t.rsplit(" ", 1)[-1])
    return t


def action_claire(l):
    a = list(dict.fromkeys(_clair(x) for x in l.get("actions") or []))
    if not a:
        return "Aucune action" if l["niveau"] == "ok" else "Surveiller"
    t = a[0]
    return t[0].upper() + t[1:] + (" (et %d autre(s) action(s))" % (len(a) - 1) if len(a) > 1 else "")


def _largeurs(ws, largeurs):
    for i, w in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _entete(ws, ligne, valeurs):
    for i, v in enumerate(valeurs, 1):
        c = ws.cell(row=ligne, column=i, value=v)
        c.font, c.fill = GRAS, ENTETE
        c.alignment = Alignment(wrap_text=True, vertical="top")


TEINTES_PHASE = {"1": "F4B6B6", "2": "F9C98A", "3": "F9E3A8", "4": "CFE2F3", "5": "E0E0E0"}


def _feuille_planning(wb, pl, position=1):
    """#878 : planning estimé en Gantt -- phases, groupes de dépendances, actions ; une colonne par semaine."""
    ws = wb.create_sheet("Planning", position)
    hy, tot = pl["hypotheses"], pl["total"]
    ws["A1"] = "Planning estimé des actions"
    ws["A1"].font = TITRE
    ws["A2"] = ("Estimation action par action, tests compris : poids en points selon un barème (%s h par point), multiplié par %s pour la "
                "marge ; %s personne(s), %s h de travail par jour, du lundi au vendredi, à partir du %s. Un serveur et les systèmes qui en "
                "dépendent sont traités ensemble." % (_n(hy.get("heures_par_point", 1)), _n(hy["marge"]), hy["personnes"], hy["heures_jour"], _fin(hy["debut"])))
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 32
    ws["A3"] = "Total : %s jours de travail (%s h sans marge, %s h avec marge) -- fin estimée le %s" % (
        str(tot["jours"]).replace(".", ","), str(tot["brut"]).replace(".", ","), str(tot["heures"]).replace(".", ","), _fin(tot["fin"]))
    ws["A3"].font = GRAS
    if not pl["taches"]:
        ws["A5"] = "Aucune action à planifier."
        return
    d0 = dt.date.fromisoformat(min(t["debut"] for t in pl["taches"]))
    d1 = dt.date.fromisoformat(max(t["fin"] for t in pl["taches"]))
    lundi = d0 - dt.timedelta(days=d0.weekday())
    semaines = []
    while lundi <= d1:
        semaines.append(lundi)
        lundi += dt.timedelta(days=7)
    cols = ["Phase / regroupement / système", "Action", "Sans marge (h)", "Avec marge (h)", "Jours", "Début", "Fin"]
    _entete(ws, 5, cols + ["sem. %s" % s_.strftime("%d/%m") for s_ in semaines])
    i = 6

    def barre(ligne, debut, fin, teinte, gras=False):
        a, b = dt.date.fromisoformat(debut), dt.date.fromisoformat(fin)
        for j, s_ in enumerate(semaines):
            if s_ <= b and s_ + dt.timedelta(days=6) >= a:
                c = ws.cell(row=ligne, column=len(cols) + 1 + j)
                c.fill = PatternFill("solid", fgColor=teinte if not gras else "888888")
    for ph in pl["phases"]:
        ts = [t for t in pl["taches"] if t["phase"] == ph["phase"]]
        for j, v in enumerate(["Phase %s : %s" % (ph["phase"], ph["nom"]), "%d système(s), %d regroupement(s)" % (ph["systemes"], ph["groupes"]),
                               round(sum(t["brut"] for t in ts), 1), ph["heures"], ph["jours"], _fin(ph["debut"]), _fin(ph["fin"])], 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.fill = GRAS, PatternFill("solid", fgColor=TEINTES_PHASE[ph["phase"]])
        barre(i, ph["debut"], ph["fin"], TEINTES_PHASE[ph["phase"]], gras=True)
        i += 1
        for g in dict.fromkeys(t["groupe"] for t in ts):
            tg = [t for t in ts if t["groupe"] == g]
            if len({t["key"] for t in tg}) > 1 or len(tg) > 1:
                ws.cell(row=i, column=1, value="  " + g).font = Font(italic=True, bold=True)
                ws.cell(row=i, column=4, value=round(sum(t["heures"] for t in tg), 1))
                ws.cell(row=i, column=5, value=round(sum(t["jours"] for t in tg), 1))
                ws.cell(row=i, column=6, value=_fin(min(t["debut"] for t in tg)))
                ws.cell(row=i, column=7, value=_fin(max(t["fin"] for t in tg)))
                i += 1
            for t in tg:
                a = _clair(t["action"])
                vals = ["      " + t["systeme"], a[0].upper() + a[1:], t["brut"], t["heures"], t["jours"], _fin(t["debut"]), _fin(t["fin"])]
                for j, v in enumerate(vals, 1):
                    ws.cell(row=i, column=j, value=v).alignment = Alignment(vertical="top", wrap_text=j <= 2)
                barre(i, t["debut"], t["fin"], TEINTES_PHASE[ph["phase"]])
                i += 1
    ws.freeze_panes = ws.cell(row=6, column=len(cols) + 1)
    _largeurs(ws, [46, 56, 10, 10, 7, 11, 11] + [7] * len(semaines))


def classeur(etat, tendance=None, organisation="", maintenant=None, plan=None):
    """-> octets du fichier .xlsx. `plan` (#878) : etat_plan.plan(etat) -- temps estimés et planning en Gantt."""
    now = maintenant or etat.get("at")
    actifs = [l for l in etat["systemes"] if l["niveau"] != "inactif" and not l.get("variante_de")]   # #879 : sans les copies
    r = etat["resume"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    ws["A1"] = "État des lieux des systèmes informatiques%s" % (" -- %s" % organisation if organisation else "")
    ws["A1"].font = TITRE
    ws["A2"] = "Situation au %s" % _jour(now)
    ws["A4"], ws["A4"].font = "Systèmes en service", GRAS
    lignes = [("Urgent : à traiter sans attendre", r.get("urgent", 0), "urgent"), ("À planifier dans les prochains mois", r.get("à planifier", 0), "à planifier"),
              ("À surveiller", r.get("à surveiller", 0), "à surveiller"), ("À jour", r.get("ok", 0), "ok"), ("Arrêtés (non utilisés)", r.get("inactif", 0), "inactif")]
    for i, (lib, n, niv) in enumerate(lignes, 5):
        ws.cell(row=i, column=1, value=lib).fill = PatternFill("solid", fgColor=TEINTES[niv])
        ws.cell(row=i, column=2, value=n)
    ws["A11"], ws["A11"].font = "Points clés", GRAS
    cles = []
    if r.get("obsoletes"):
        cles.append("%d système(s) en service ne reçoivent plus de correctifs de sécurité de leur éditeur." % r["obsoletes"])
    if r.get("secu"):
        cles.append("%d système(s) ont des correctifs de sécurité disponibles mais non appliqués." % r["secu"])
    if r.get("exposes"):
        cles.append("%d système(s) sont accessibles depuis Internet." % r["exposes"])
    if r.get("decouverts"):   # #873
        cles.append("%d appareil(s) vus sur le réseau ne sont suivis par aucun outil : à identifier." % r["decouverts"])
    dep = sum(1 for l in actifs if len(l["dependants"]) >= 2)
    if dep:
        cles.append("%d système(s) sont utilisés par plusieurs autres : une panne aurait des effets en chaîne." % dep)
    if plan and plan["taches"]:   # #878
        cles.append("Temps de travail estimé : %s jours (marge ×%s comprise, tests inclus), fin estimée le %s avec %s personne(s) -- voir « Planning »."
                    % (_n(plan["total"]["jours"]), _n(plan["hypotheses"]["marge"]), _fin(plan["total"]["fin"]), plan["hypotheses"]["personnes"]))
    if not cles:
        cles.append("Aucun point critique relevé.")
    for i, t in enumerate(cles, 12):
        ws.cell(row=i, column=1, value="• " + t)
    base = 13 + len(cles)
    ws.cell(row=base, column=1, value="Priorités").font = GRAS
    _entete(ws, base + 1, ["Système", "Pourquoi", "Action proposée", "Niveau"])
    prio = [l for l in actifs if l["niveau"] in ("urgent", "à planifier")][:10]
    for i, l in enumerate(prio, base + 2):
        ws.cell(row=i, column=1, value=nom(l))
        ws.cell(row=i, column=2, value=raisons_claires(l)).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=3, value=action_claire(l)).alignment = Alignment(wrap_text=True, vertical="top")
        c = ws.cell(row=i, column=4, value=LIBELLE_NIVEAU[l["niveau"]])
        c.fill = PatternFill("solid", fgColor=TEINTES[l["niveau"]])
    if not prio:
        ws.cell(row=base + 2, column=1, value="Aucune priorité.")
    _largeurs(ws, [48, 60, 50, 14])

    ws2 = wb.create_sheet("Systèmes")
    cols = ["Système", "Nature", "Rôle", "Système d'exploitation", "Fin des correctifs", "Dernière mise à jour", "Correctifs de sécurité en attente",
            "Utilisation", "Accessible depuis Internet", "Utilisé par (nombre de systèmes)", "Niveau", "Pourquoi", "Action proposée",
            "Temps estimé (jours, marge comprise)"]
    _entete(ws2, 1, cols)
    import etat_plan as _ep   # #878
    temps = _ep.par_systeme(plan) if plan else {}
    for i, l in enumerate(etat["systemes"], 2):
        if l.get("variante_de"):   # #879 : copie / clone / test : listé pour la forme, sans niveau ni action
            for j, v in enumerate([nom(l), TYPE.get(l["type"], l["type"]), "—", l["os"]["nom"] or "—", "", "", "", "", "", "", "Copie / essai",
                                   "rattaché à « %s » (%s)" % (l["variante_de"]["name"], l["variante_de"]["role"]), "aucune (suivi avec sa référence)", "—"], 1):
                ws2.cell(row=i, column=j, value=v).font = Font(italic=True, color="777777")
            continue
        vals = [nom(l), TYPE.get(l["type"], l["type"]), ", ".join(l["roles"]) or "—", l["os"]["nom"] or "non relevé", _fin(l["os"]["fin"]),
                _jour(l["maj_at"]), l["attente_secu"] if l["attente_secu"] is not None else "—", USAGE.get(l["usage"], l["usage"]),
                "Oui" if l["expose"] else "Non", len(l["dependants"]), LIBELLE_NIVEAU[l["niveau"]], raisons_claires(l), action_claire(l),
                (temps.get(l["key"]) or {}).get("jours", "—")]
        for j, v in enumerate(vals, 1):
            c = ws2.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws2.cell(row=i, column=11).fill = PatternFill("solid", fgColor=TEINTES[l["niveau"]])
    ws2.freeze_panes = "B2"

    # #873 : audit réseau -- une ligne par service vu sur le réseau (suivi de migration : qui expose quoi, en quelle version)
    vus = [l for l in etat["systemes"] if l.get("reseau")]
    if vus:
        ws3 = wb.create_sheet("Audit réseau")
        _entete(ws3, 1, ["Système", "Adresse", "Nature", "Système d'exploitation", "Port", "Service", "Logiciel et version", "Source", "Vu le"])
        i = 2
        for l in vus:
            for p in l["reseau"]:
                vals = [nom(l), ", ".join(l["ips"]), TYPE.get(l["type"], l["type"]), l["os"]["nom"] or "non identifié", p["port"], p.get("service") or "—",
                        p.get("logiciel") or "—", l["source"], (l.get("vu_le") or "")[:10] or "—"]
                for j, v in enumerate(vals, 1):
                    ws3.cell(row=i, column=j, value=v)
                i += 1
        ws3.freeze_panes = "B2"
        ws3.auto_filter.ref = "A1:I%d" % max(1, i - 1)
        _largeurs(ws3, [30, 16, 12, 34, 7, 14, 40, 14, 11])
    ws2.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), max(1, len(etat["systemes"]) + 1))
    _largeurs(ws2, [34, 16, 30, 26, 14, 14, 12, 12, 12, 12, 13, 50, 44, 14])

    ch = etat.get("chantiers") or {}
    if ch.get("decisions") or ch.get("chantiers") or ch.get("migration") or ch.get("organisation") or ch.get("infrastructure"):
        wsc = wb.create_sheet("Chantiers", 1)
        wsc["A1"] = "Chantiers en cours et décisions à prendre"
        wsc["A1"].font = TITRE
        ligne = 3
        if ch.get("decisions"):
            wsc.cell(row=ligne, column=1, value="Décisions à prendre").font = GRAS
            ligne += 1
            _entete(wsc, ligne, ["Objet", "Pourquoi", "Montant", "Échéance"])
            for d in ch["decisions"]:
                ligne += 1
                for j, v in enumerate([d.get("titre"), d.get("detail"), d.get("montant"), d.get("echeance")], 1):
                    wsc.cell(row=ligne, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
            ligne += 2
        if ch.get("infrastructure"):
            wsc.cell(row=ligne, column=1, value="Serveurs et sites concernés").font = GRAS
            ligne += 1
            _entete(wsc, ligne, ["Serveur", "Site", "Machines hébergées", "Capacité (To)", "Occupé", "Disque de secours", "État"])
            for x in ch["infrastructure"]:
                ligne += 1
                vals = [x["nom"] + (" (dans %s)" % x["dans"] if x.get("dans") else ""), x["site"], x["machines"], x["espace_to"],
                        ("%s %%" % x["occupe_pct"]) if x["occupe_pct"] is not None else "—", x["secours"], x["etat"]]
                for j, v in enumerate(vals, 1):
                    c = wsc.cell(row=ligne, column=j, value=v)
                    if (j == 6 and v == "non") or (j == 7 and v != "normal"):
                        c.fill = PatternFill("solid", fgColor=TEINTES["urgent"])
            ligne += 2
        if ch.get("organisation"):
            wsc.cell(row=ligne, column=1, value="Organisation du chantier : où, quand, combien").font = GRAS
            ligne += 1
            _entete(wsc, ligne, ["Phase", "Quand", "Où (de → vers)", "Volume (Go)", "Durée estimée", "Coupure de service", "Contraintes", "Avancement"])
            for x in ch["organisation"]:
                ligne += 1
                vals = [x["titre"], x["quand"] or "à fixer", x["ou"] or "—", x["volume_go"] or "—", x["duree"] or "—", x["coupure"] or "—",
                        x["contraintes"] or "—", "%s (%d / %d)" % (x["etat"], x["faites"], x["total"])]
                for j, v in enumerate(vals, 1):
                    wsc.cell(row=ligne, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
            ligne += 2
        m = ch.get("migration")
        if m:
            wsc.cell(row=ligne, column=1, value="%s : %d étape(s) faite(s) sur %d" % (m["nom"], m["faites"], m["total"])).font = GRAS
            for p in m["phases"]:
                ligne += 1
                wsc.cell(row=ligne, column=1, value=p["nom"])
                wsc.cell(row=ligne, column=2, value="%d / %d" % (p["faites"], p["total"]))
            ligne += 2
        for c in ch.get("chantiers") or []:
            wsc.cell(row=ligne, column=1, value=c["chantier"]).font = GRAS
            ligne += 1
            _entete(wsc, ligne, ["Où en est-on", "Date", "Action", "Résultat / détail"])
            for lib, k in (("Fait (%d derniers jours)" % ch.get("periode_jours", 45), "faits"), ("En cours", "en_cours"), ("Planifié", "planifies"), ("À faire", "a_faire")):
                for e in c[k]:
                    ligne += 1
                    for j, v in enumerate([lib, _fin(e.get("day")), e.get("acte"), e.get("resultat") or e.get("description")], 1):
                        wsc.cell(row=ligne, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
            ligne += 2
        _largeurs(wsc, [40, 16, 30, 12, 14, 18, 40, 18])

    ws3 = wb.create_sheet("Dépendances")
    _entete(ws3, 1, ["Système", "utilise", "pour"])
    noms = {l["key"]: nom(l) for l in etat["systemes"]}
    for i, d in enumerate(etat.get("dependances") or [], 2):
        ws3.cell(row=i, column=1, value=noms.get(d["de"], d["de"]))
        ws3.cell(row=i, column=2, value=noms.get(d["vers"], d["vers"]))
        ws3.cell(row=i, column=3, value=d["nature"])
    if not etat.get("dependances"):
        ws3.cell(row=2, column=1, value="Aucune dépendance relevée entre les systèmes suivis.")
    _largeurs(ws3, [40, 40, 30])

    if plan is not None:   # #878 : juste après la synthèse
        _feuille_planning(wb, plan, 1)
    ws4 = wb.create_sheet("Évolution")
    _entete(ws4, 1, ["Jour", "Urgent", "À planifier", "À surveiller", "À jour", "Arrêtés", "Plus maintenus", "Correctifs en attente"])
    for i, t in enumerate((tendance or [])[-120:], 2):
        for j, k in enumerate(["jour", "urgent", "à planifier", "à surveiller", "ok", "inactif", "obsoletes", "secu"], 1):
            ws4.cell(row=i, column=j, value=_fin(t.get(k)) if k == "jour" else t.get(k))
    if not tendance:
        ws4.cell(row=2, column=1, value="Le suivi commence aujourd'hui : une ligne est ajoutée chaque jour.")
    _largeurs(ws4, [14, 10, 12, 12, 10, 10, 14, 18])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
